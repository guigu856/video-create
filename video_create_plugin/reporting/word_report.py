"""将结构化参考学习报告确定性排版为单份人类可读 DOCX。"""

from __future__ import annotations

import zipfile
from collections.abc import Iterable
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from .models import EvidenceBackedClaim, ReferenceStudyReport

_FONT = "Microsoft YaHei"
_BLUE = "2E74B5"
_DARK_BLUE = "1F4D78"
_NAVY = "203748"
_MUTED = "666666"
_TABLE_FILL = "E8EEF5"
_TABLE_BORDER = "B7C6D6"
_CONTENT_WIDTH_DXA = 9360
_TABLE_INDENT_DXA = 120
_CELL_MARGIN_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
_FIXED_TIME = datetime(2000, 1, 1, tzinfo=UTC)
_ZIP_TIME = (2000, 1, 1, 0, 0, 0)

_OVERVIEW_LABELS = {
    "video_type_and_core_mechanism": "视频类型与核心机制",
    "production_method": "整体制作方法",
    "visual_language": "视觉语言与画面组织",
    "rhythm_and_sound": "节奏与声音使用方法",
    "transition_principles": "转场与镜头连接原则",
    "asset_and_music_traits": "素材与音乐的性质",
    "viewing_experience": "预期观看体验",
}
_AUDIO_SCOPE_LABELS = {
    "mixed_program_audio": "成片混合音轨",
    "isolated_bgm": "独立背景音乐",
}
_FACT_STATUS_LABELS = {
    "measured": "测量事实",
    "algorithm_candidate": "算法候选",
    "agent_inference": "Agent 推断",
    "reconstruction_suggestion": "重构建议",
}
_STAGE_LABELS = {
    "stage1": "阶段一：创意方向",
    "stage2": "阶段二：素材与音乐准备",
    "stage3": "阶段三：剪辑规格",
}


def write_human_report(report: ReferenceStudyReport, path: Path) -> None:
    """写入字节稳定的单份 DOCX 报告。"""
    document = Document()
    _configure_document(document)
    _add_cover(document, report)
    _add_page_break(document)
    _add_overview(document, report)
    _add_bgm_analysis(document, report)
    _add_shot_analysis(document, report)
    _add_editing_grammar(document, report)
    _add_creation_projection(document, report)
    _add_evidence_index(document, report)
    _save_deterministically(document, path)


def _configure_document(document: DocumentType) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    _set_style(styles["Normal"], 11, "000000", 0, 6, 1.25)
    _set_style(styles["Title"], 28, _NAVY, 0, 8, 1.0)
    _set_style(styles["Subtitle"], 14, _DARK_BLUE, 0, 18, 1.0)
    _set_style(styles["Heading 1"], 16, _BLUE, 18, 10, 1.0)
    _set_style(styles["Heading 2"], 13, _BLUE, 14, 7, 1.0)
    _set_style(styles["Heading 3"], 12, _DARK_BLUE, 10, 5, 1.0)

    core = document.core_properties
    core.title = "参考视频分析报告"
    core.subject = "视觉、声音、音画关系与可复用剪辑语法"
    core.author = "video-create"
    core.last_modified_by = "video-create"
    core.created = _FIXED_TIME
    core.modified = _FIXED_TIME
    core.revision = 1

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_run(header, "参考视频分析报告", size=9, color=_MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(footer, "第 ", size=9, color=_MUTED)
    _add_page_field(footer)
    _add_run(footer, " 页", size=9, color=_MUTED)


def _set_style(
    style: Any,
    size: float,
    color: str,
    before: float,
    after: float,
    line_spacing: float,
) -> None:
    font = style.font
    font.name = _FONT
    font.size = Pt(size)
    font.color.rgb = RGBColor.from_string(color)
    _set_font_mapping(style.element.get_or_add_rPr())
    paragraph_format = style.paragraph_format
    paragraph_format.space_before = Pt(before)
    paragraph_format.space_after = Pt(after)
    paragraph_format.line_spacing = line_spacing
    paragraph_format.widow_control = True


def _set_font_mapping(run_properties: Any) -> None:
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), _FONT)


def _add_cover(document: DocumentType, report: ReferenceStudyReport) -> None:
    for _ in range(3):
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(18)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    _add_run(kicker, "REFERENCE STUDY", size=10, color=_BLUE, bold=True)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(title, "参考视频分析报告", size=28, color=_NAVY, bold=True)

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(
        subtitle,
        "视觉、声音、音画关系与可复用剪辑语法",
        size=14,
        color=_DARK_BLUE,
    )

    metadata = (
        ("报告 ID", report.report_id),
        ("分析 ID", report.analysis_id),
        ("视频时长", _duration(report.duration_us)),
        ("分析版本", report.analysis_version),
        ("媒体哈希", report.source_media_sha256),
    )
    for label, value in metadata:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)
        _add_run(paragraph, f"{label}：", size=9.5, color=_DARK_BLUE, bold=True)
        _add_run(paragraph, value, size=9.5, color=_MUTED)


def _add_overview(document: DocumentType, report: ReferenceStudyReport) -> None:
    document.add_heading("参考视频总体理解", level=1)
    for section in report.video_overview.sections:
        document.add_heading(_OVERVIEW_LABELS[section.section], level=2)
        _add_claims(document, section.claims)


def _add_bgm_analysis(document: DocumentType, report: ReferenceStudyReport) -> None:
    bgm = report.bgm_analysis
    document.add_heading("BGM 与音画关系分析", level=1)
    _add_labeled_paragraph(document, "音频范围", _AUDIO_SCOPE_LABELS[bgm.audio_scope])

    document.add_heading("音乐段落", level=2)
    for section in bgm.sections:
        document.add_heading(
            f"{section.section_id} · {section.section_type} · "
            f"{_time_range(section.time_range.start_us, section.time_range.end_us)}",
            level=3,
        )
        _add_claims(document, section.claims)

    document.add_heading("Tempo 候选", level=2)
    if bgm.tempo_candidates:
        table = document.add_table(rows=1, cols=3)
        _set_header_row(table, ("BPM", "置信度", "证据"))
        for item in bgm.tempo_candidates:
            cells = table.add_row().cells
            _set_cell_text(cells[0], f"{item.bpm:g}")
            _set_cell_text(cells[1], f"{item.confidence:.2f}")
            _set_cell_text(cells[2], "、".join(item.evidence_refs))
        _format_table(table, (1400, 1600, 6360))
    else:
        _add_labeled_paragraph(document, "结果", "当前报告未记录 Tempo 候选")

    document.add_heading("音画关系", level=2)
    for relation in bgm.audiovisual_relations:
        document.add_heading(
            f"{relation.relation_type} · "
            f"{_time_range(relation.time_range.start_us, relation.time_range.end_us)}",
            level=3,
        )
        _add_claims(document, (relation.claim,))


def _add_shot_analysis(document: DocumentType, report: ReferenceStudyReport) -> None:
    document.add_heading("逐镜分析", level=1)
    _add_labeled_paragraph(
        document,
        "时间线覆盖",
        f"0.000s–{report.duration_us / 1_000_000:.3f}s，共 "
        f"{len(report.shot_analysis.shots)} 个主镜头",
    )
    for shot in report.shot_analysis.shots:
        document.add_heading(
            f"镜头 {shot.shot_id} · "
            f"{_time_range(shot.time_range.start_us, shot.time_range.end_us)}",
            level=2,
        )
        fields = (
            ("节奏单元", shot.rhythm_unit_id),
            ("主画面", shot.main_picture),
            ("可见图层", "、".join(shot.visible_layers)),
            ("画面变化", shot.visual_change),
            ("效果作用域", shot.effect_scope),
            ("声音关系", shot.sound_relation),
            ("镜头连接", shot.connection),
            ("剪辑句子功能", shot.sentence_function),
        )
        for label, value in fields:
            _add_labeled_paragraph(document, label, value)
        document.add_heading("证据结论", level=3)
        _add_claims(document, shot.claims)


def _add_editing_grammar(document: DocumentType, report: ReferenceStudyReport) -> None:
    grammar = report.editing_grammar
    document.add_heading("可复用剪辑语法", level=1)
    document.add_heading("核心观看体验", level=2)
    _add_claims(document, grammar.viewing_experience)

    document.add_heading("可迁移剪辑句子", level=2)
    for index, sentence in enumerate(grammar.reusable_sentences, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        _add_run(paragraph, f"剪辑句子 {index}　", bold=True, color=_DARK_BLUE)
        _add_run(paragraph, sentence.content)
        _add_metadata_line(
            document,
            f"ID：{sentence.sentence_id}　置信度：{sentence.confidence:.2f}　"
            f"证据：{'、'.join(sentence.evidence_refs)}",
        )

    document.add_heading("原片专属特征", level=2)
    _add_claims(document, grammar.reference_specific_notes)


def _add_creation_projection(document: DocumentType, report: ReferenceStudyReport) -> None:
    document.add_heading("创作阶段知识投影", level=1)
    projection = report.creation_context_projection
    for stage, items in (
        ("stage1", projection.stage1_knowledge),
        ("stage2", projection.stage2_knowledge),
        ("stage3", projection.stage3_knowledge),
    ):
        document.add_heading(_STAGE_LABELS[stage], level=2)
        if not items:
            _add_labeled_paragraph(document, "结果", "当前阶段没有投影知识")
            continue
        table = document.add_table(rows=1, cols=4)
        _set_header_row(table, ("知识类型", "内容", "置信度", "证据"))
        for item in items:
            cells = table.add_row().cells
            _set_cell_text(cells[0], item.knowledge_type)
            _set_cell_text(cells[1], item.content)
            _set_cell_text(cells[2], f"{item.confidence:.2f}")
            _set_cell_text(cells[3], "、".join(item.evidence_refs))
        _format_table(table, (1800, 3960, 1200, 2400))


def _add_evidence_index(document: DocumentType, report: ReferenceStudyReport) -> None:
    document.add_heading("证据索引", level=1)
    _add_labeled_paragraph(
        document,
        "说明",
        "报告中的证据 ID 可在 EvidenceBundle 中解析为工作区文件及其 SHA-256。",
    )
    table = document.add_table(rows=1, cols=4)
    _set_header_row(table, ("证据 ID", "类型", "时间范围", "文件"))
    for entry in report.evidence_bundle.entries:
        cells = table.add_row().cells
        _set_cell_text(cells[0], entry.evidence_id)
        _set_cell_text(cells[1], entry.evidence_type)
        _set_cell_text(
            cells[2],
            (
                _time_range(entry.time_range.start_us, entry.time_range.end_us)
                if entry.time_range
                else "全局"
            ),
        )
        _set_cell_text(cells[3], entry.file.path)
    _format_table(table, (2200, 1900, 1800, 3460))


def _add_claims(
    document: DocumentType,
    claims: Iterable[EvidenceBackedClaim],
) -> None:
    for index, claim in enumerate(claims, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        _add_run(paragraph, f"结论 {index}　", bold=True, color=_DARK_BLUE)
        _add_run(paragraph, claim.content)
        evidence = "、".join(claim.evidence_refs) or "无"
        _add_metadata_line(
            document,
            f"{_FACT_STATUS_LABELS[claim.fact_status]}　置信度 {claim.confidence:.2f}　"
            f"证据 {evidence}",
        )


def _add_labeled_paragraph(document: DocumentType, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    _add_run(paragraph, f"{label}：", bold=True, color=_DARK_BLUE)
    _add_run(paragraph, value)


def _add_metadata_line(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True
    _add_run(paragraph, text, size=9, color=_MUTED)


def _add_run(
    paragraph: Paragraph,
    text: str,
    *,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
) -> None:
    run = paragraph.add_run(text)
    run.font.name = _FONT
    _set_font_mapping(run._element.get_or_add_rPr())
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _set_cell_text(
    cell: _Cell,
    text: str,
    *,
    bold: bool = False,
    color: str = "000000",
) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    _add_run(paragraph, text, size=9.5, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_header_row(table: Table, labels: tuple[str, ...]) -> None:
    row = table.rows[0]
    for cell, label in zip(row.cells, labels):
        _set_cell_text(cell, label, bold=True, color=_DARK_BLUE)
        _shade_cell(cell, _TABLE_FILL)
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:tblHeader"))


def _format_table(
    table: Table,
    widths_dxa: tuple[int, ...],
) -> None:
    if sum(widths_dxa) != _CONTENT_WIDTH_DXA:
        raise ValueError("DOCX 表格列宽总和必须为 9360 DXA")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    properties = table._tbl.tblPr
    _set_measure(properties, "tblW", _CONTENT_WIDTH_DXA)
    _set_measure(properties, "tblInd", _TABLE_INDENT_DXA)

    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    _set_table_borders(properties)
    _set_cell_margins(properties)

    grid = table._tbl.tblGrid
    for child in tuple(grid):
        grid.remove(child)
    for width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            _set_measure(cell._tc.get_or_add_tcPr(), "tcW", width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_measure(parent: Any, tag: str, value: int) -> None:
    element = parent.find(qn(f"w:{tag}"))
    if element is None:
        element = OxmlElement(f"w:{tag}")
        parent.append(element)
    element.set(qn("w:w"), str(value))
    element.set(qn("w:type"), "dxa")


def _set_cell_margins(properties: Any) -> None:
    margins = properties.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        properties.append(margins)
    for side, value in _CELL_MARGIN_DXA.items():
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _set_table_borders(properties: Any) -> None:
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), _TABLE_BORDER)


def _shade_cell(cell: _Cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _add_page_field(paragraph: Paragraph) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    _set_font_mapping(properties)
    run.append(properties)
    run.extend((begin, instruction, separate, text, end))
    paragraph._p.append(run)


def _add_page_break(document: DocumentType) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    page_break = OxmlElement("w:br")
    page_break.set(qn("w:type"), "page")
    run._r.append(page_break)


def _duration(duration_us: int) -> str:
    total_seconds = duration_us / 1_000_000
    minutes, seconds = divmod(total_seconds, 60)
    return f"{int(minutes):02d}:{seconds:06.3f}"


def _time_range(start_us: int, end_us: int) -> str:
    return f"{start_us / 1_000_000:.3f}s–{end_us / 1_000_000:.3f}s"


def _save_deterministically(document: DocumentType, path: Path) -> None:
    raw_path = path.with_name(f"{path.stem}.raw.docx")
    temporary_path = path.with_name(f"{path.stem}.tmp.docx")
    document.save(str(raw_path))
    try:
        with (
            zipfile.ZipFile(raw_path, "r") as source,
            zipfile.ZipFile(
                temporary_path,
                "w",
                compression=zipfile.ZIP_DEFLATED,
                compresslevel=9,
            ) as target,
        ):
            for name in sorted(source.namelist()):
                original = source.getinfo(name)
                info = zipfile.ZipInfo(name, _ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = original.external_attr
                info.create_system = original.create_system
                target.writestr(info, source.read(name))
        temporary_path.replace(path)
    finally:
        raw_path.unlink(missing_ok=True)
        temporary_path.unlink(missing_ok=True)
